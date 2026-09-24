"""PDF / Excel dışa aktarım — etüt hocası indirmeleri."""

from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _excel_response(workbook: Workbook, filename: str) -> HttpResponse:
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _pdf_response(buffer: BytesIO, filename: str) -> HttpResponse:
    response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _styles():
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "CSTitle",
        parent=styles["Heading1"],
        fontSize=14,
        spaceAfter=8,
        textColor=colors.HexColor("#0c1220"),
    )
    sub = ParagraphStyle(
        "CSSub",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#4a5568"),
        spaceAfter=12,
    )
    return title, sub


def _table(data, col_widths=None):
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1557a0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d7dde8")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f9")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def excel_siralama(etut, deneme, siralama: list[dict]) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Siralama"
    ws.append(["Sira", "Talebe", "Sinif", "Ortalama %", "Konu sayisi"])
    for row in siralama:
        ws.append(
            [
                row["sira"],
                row["ad_soyad"],
                row["sinif"],
                float(row["ortalama"]) if row["ortalama"] is not None else None,
                row["konu_sayisi"],
            ]
        )
    name = f"siralama_{etut.id}_{deneme.id}.xlsx"
    return _excel_response(wb, name)


def excel_kazanim(etut, deneme, kazanimlar: list[dict]) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kazanımlar"
    ws.append(["Ders", "Kazanım", "Etüt ort. %", "Katılan", "Zayıf"])
    for row in kazanimlar:
        ws.append(
            [
                row["ders"],
                row["konu"],
                float(row["ortalama"]) if row["ortalama"] is not None else None,
                row["talebe_sayisi"],
                "Evet" if row.get("zayif") else "Hayır",
            ]
        )
    name = f"kazanim_{etut.id}_{deneme.id}.xlsx"
    return _excel_response(wb, name)


def excel_sinif_raporu(deneme, sinif, rapor: list[dict]) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sinif Raporu"
    ws.append(["Ders", "Kazanım", "Ortalama %", "Talebe"])
    for row in rapor:
        ws.append(
            [
                row["ders"],
                row["konu"],
                float(row["ortalama"]) if row["ortalama"] is not None else None,
                row["talebe_sayisi"],
            ]
        )
    label = sinif.ad if sinif else "tum"
    name = f"sinif_raporu_{deneme.id}_{label}.xlsx"
    return _excel_response(wb, name)


def excel_talebe_kazanim(talebe, kutular: list[dict]) -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = "Talebe Kazanım"
    ws.append(["Deneme", "Tarih", "Ders", "Kazanım", "Yüzde", "Net", "Zayıf"])
    for kutu in kutular:
        for row in kutu["kazanimlar"]:
            net = ""
            if row["net_dogru"] is not None and row["net_toplam"] is not None:
                net = f"{row['net_dogru']}/{row['net_toplam']}"
            ws.append(
                [
                    kutu["deneme"].ad,
                    kutu["deneme"].tarih.isoformat(),
                    row["ders"],
                    row["konu"],
                    float(row["yuzde"]) if row["yuzde"] is not None else None,
                    net,
                    "Evet" if row["zayif"] else "Hayır",
                ]
            )
    safe = talebe.ad_soyad.replace(" ", "_")[:40]
    return _excel_response(wb, f"talebe_kazanim_{safe}.xlsx")


def pdf_siralama(etut, deneme, siralama: list[dict]) -> HttpResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm)
    title, sub = _styles()
    story = [
        Paragraph(f"Genel Sıralama — {deneme.ad}", title),
        Paragraph(f"{etut.ad} · {deneme.tarih.strftime('%d.%m.%Y')}", sub),
    ]
    data = [["#", "Talebe", "Sınıf", "Ort. %", "Konu"]]
    for row in siralama:
        data.append(
            [
                str(row["sira"]),
                row["ad_soyad"],
                row["sinif"],
                f"{row['ortalama']}%" if row["ortalama"] is not None else "—",
                str(row["konu_sayisi"]),
            ]
        )
    story.append(_table(data, col_widths=[12 * mm, 70 * mm, 25 * mm, 25 * mm, 20 * mm]))
    doc.build(story)
    buffer.seek(0)
    return _pdf_response(buffer, f"siralama_{etut.id}_{deneme.id}.pdf")


def pdf_kazanim(etut, deneme, kazanimlar: list[dict]) -> HttpResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
    )
    title, sub = _styles()
    story = [
        Paragraph(f"Kazanım Listesi — {deneme.ad}", title),
        Paragraph(
            f"{etut.ad} · {deneme.tarih.strftime('%d.%m.%Y')} · %70 altı zayıf",
            sub,
        ),
    ]
    data = [["Ders", "Kazanım", "Etüt ort.", "Katılan", "Durum"]]
    for row in kazanimlar:
        data.append(
            [
                row["ders"][:28],
                row["konu"][:42],
                f"{row['ortalama']}%" if row["ortalama"] is not None else "—",
                str(row["talebe_sayisi"]),
                "ZAYIF" if row.get("zayif") else "OK",
            ]
        )
    story.append(
        _table(data, col_widths=[45 * mm, 90 * mm, 25 * mm, 20 * mm, 20 * mm])
    )
    doc.build(story)
    buffer.seek(0)
    return _pdf_response(buffer, f"kazanim_{etut.id}_{deneme.id}.pdf")


def pdf_sinif_raporu(deneme, sinif, rapor: list[dict]) -> HttpResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
    )
    title, sub = _styles()
    sinif_ad = sinif.ad if sinif else "Tüm sınıflar"
    story = [
        Paragraph(f"Sınıf Kazanım Raporu — {deneme.ad}", title),
        Paragraph(f"{sinif_ad} · {deneme.tarih.strftime('%d.%m.%Y')}", sub),
    ]
    data = [["Ders", "Kazanım", "Ort. %", "Talebe"]]
    for row in rapor:
        data.append(
            [
                row["ders"][:28],
                row["konu"][:48],
                f"{row['ortalama']}%" if row["ortalama"] is not None else "—",
                str(row["talebe_sayisi"]),
            ]
        )
    story.append(_table(data, col_widths=[50 * mm, 120 * mm, 25 * mm, 20 * mm]))
    doc.build(story)
    buffer.seek(0)
    label = sinif.ad if sinif else "tum"
    return _pdf_response(buffer, f"sinif_raporu_{deneme.id}_{label}.pdf")


def pdf_talebe_kazanim(etut, talebe, kutular: list[dict]) -> HttpResponse:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
    )
    title, sub = _styles()
    story = [
        Paragraph(f"Talebe Kazanım Listesi — {talebe.ad_soyad}", title),
        Paragraph(
            f"{etut.ad} · {talebe.sinif.ad} · Zayıf konular (%70 altı) işaretli",
            sub,
        ),
    ]

    for kutu in kutular:
        story.append(
            Paragraph(
                f"{kutu['deneme'].ad} ({kutu['deneme'].tarih.strftime('%d.%m.%Y')}) "
                f"— Ort: {kutu['ortalama'] if kutu['ortalama'] is not None else '—'}% "
                f"— Zayıf: {kutu['zayif_sayisi']}",
                ParagraphStyle(
                    "DenemeHead",
                    fontSize=10,
                    fontName="Helvetica-Bold",
                    spaceBefore=8,
                    spaceAfter=4,
                ),
            )
        )
        data = [["Ders", "Kazanım", "%", "Net", "Durum"]]
        # Zayıflar önce zaten sıralı
        for row in kutu["kazanimlar"]:
            net = "—"
            if row["net_dogru"] is not None and row["net_toplam"] is not None:
                net = f"{row['net_dogru']}/{row['net_toplam']}"
            data.append(
                [
                    row["ders"][:18],
                    row["konu"][:34],
                    f"{row['yuzde']}%" if row["yuzde"] is not None else "—",
                    net,
                    "ZAYIF" if row["zayif"] else "OK",
                ]
            )
        story.append(
            _table(data, col_widths=[32 * mm, 70 * mm, 18 * mm, 22 * mm, 18 * mm])
        )
        story.append(Spacer(1, 6))

    doc.build(story)
    buffer.seek(0)
    safe = talebe.ad_soyad.replace(" ", "_")[:40]
    return _pdf_response(buffer, f"talebe_kazanim_{safe}.pdf")
